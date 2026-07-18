"""Document reader: PDF/docx/xlsx/text extraction (+ OCR where tesseract is present)."""
import asyncio
import shutil

import pytest

from engine.tools.documents import ReadDocumentTool, extract_document
from engine.tools.files import FileWorkspace


def _pdf_with_text(path, text):
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=14)
    doc.save(path)
    doc.close()


def test_extract_text_pdf(tmp_path):
    p = str(tmp_path / "doc.pdf")
    _pdf_with_text(p, "Quarterly revenue was 1.2 million dollars")
    out = extract_document(p)
    assert "Quarterly revenue" in out and "million" in out


def test_extract_docx(tmp_path):
    import docx
    p = str(tmp_path / "d.docx")
    d = docx.Document(); d.add_paragraph("Hello from a Word document."); d.save(p)
    assert "Word document" in extract_document(p)


def test_extract_xlsx(tmp_path):
    import openpyxl
    p = str(tmp_path / "s.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["name", "score"]); ws.append(["Alice", 91]); wb.save(p)
    out = extract_document(p)
    assert "Alice" in out and "91" in out and "score" in out


def test_extract_plaintext(tmp_path):
    p = tmp_path / "n.txt"; p.write_text("just some notes")
    assert "just some notes" in extract_document(str(p))


def test_unsupported_type(tmp_path):
    p = tmp_path / "x.zip"; p.write_bytes(b"PK\x03\x04")
    with pytest.raises(ValueError):
        extract_document(str(p))


def test_read_document_tool_missing(tmp_path):
    ws = FileWorkspace(str(tmp_path / "ws"))
    t = ReadDocumentTool(ws)
    assert "no file" in asyncio.run(t.run(t.Params(name="nope.pdf"))).lower()


def test_read_document_tool_pdf(tmp_path):
    ws = FileWorkspace(str(tmp_path / "ws"))
    import os
    os.makedirs(ws.root, exist_ok=True)
    _pdf_with_text(os.path.join(ws.root, "r.pdf"), "Sales were 82 units last week")
    t = ReadDocumentTool(ws)
    assert "82" in asyncio.run(t.run(t.Params(name="r.pdf")))


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract not installed")
def test_ocr_scanned_pdf(tmp_path):
    """A PDF page with NO text layer (image only) is read via OCR."""
    import fitz
    from PIL import Image, ImageDraw
    img_path = str(tmp_path / "scan.png")
    im = Image.new("RGB", (600, 200), "white")
    ImageDraw.Draw(im).text((20, 80), "SCANNED INVOICE 12345", fill="black")
    im.save(img_path)
    p = str(tmp_path / "scan.pdf")
    doc = fitz.open(); page = doc.new_page()
    page.insert_image(fitz.Rect(0, 0, 600, 200), filename=img_path); doc.save(p); doc.close()
    out = extract_document(p)
    assert "SCANNED" in out.upper() or "12345" in out
